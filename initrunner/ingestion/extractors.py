"""File-type dispatcher for text extraction. Heavy formats use lazy imports."""

from __future__ import annotations

import csv
import io
import json
from collections.abc import Callable
from pathlib import Path

# ---------------------------------------------------------------------------
# Extractor registry
# ---------------------------------------------------------------------------

_EXTRACTORS: dict[str, Callable[[Path], str]] = {}


_ExtractorFn = Callable[[Path], str]


def register_extractor(*extensions: str) -> Callable[[_ExtractorFn], _ExtractorFn]:
    """Register a function as the extractor for one or more file extensions."""

    def decorator(fn: _ExtractorFn) -> _ExtractorFn:
        for ext in extensions:
            _EXTRACTORS[ext] = fn
        return fn

    return decorator


def extract_text(path: Path) -> str:
    """Extract text from a file based on its extension."""
    suffix = path.suffix.lower()
    extractor = _EXTRACTORS.get(suffix)
    if extractor is None:
        raise ValueError(f"Unsupported file type: {suffix}")
    return extractor(path)


# ---------------------------------------------------------------------------
# Built-in extractors (registered at import time)
# ---------------------------------------------------------------------------


@register_extractor(".txt", ".md", ".rst")
def _extract_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8")


@register_extractor(".csv")
def _extract_csv(path: Path) -> str:
    with path.open(encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)
    return "\n".join(", ".join(row) for row in rows)


@register_extractor(".json")
def _extract_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8"))
    return json.dumps(data, indent=2, ensure_ascii=False)


@register_extractor(".html", ".htm")
def _extract_html(path: Path) -> str:
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    html = path.read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "html.parser")
    # Remove script and style tags
    for tag in soup(["script", "style"]):
        tag.decompose()
    return markdownify(str(soup)).strip()


@register_extractor(".pdf")
def _extract_pdf(path: Path) -> str:
    from initrunner._compat import require_ingest

    require_ingest("pymupdf4llm")
    import pymupdf4llm  # type: ignore[unresolved-import]

    return pymupdf4llm.to_markdown(str(path))


@register_extractor(".docx")
def _extract_docx(path: Path) -> str:
    from initrunner._compat import require_ingest

    require_ingest("docx")
    import docx  # type: ignore[unresolved-import]

    doc = docx.Document(str(path))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_url(url: str, *, timeout: int = 15, max_bytes: int = 512_000) -> str:
    """Fetch a URL and extract text as markdown."""
    from initrunner._html import fetch_url_as_markdown

    return fetch_url_as_markdown(url, timeout=timeout, max_bytes=max_bytes)


@register_extractor(".xlsx")
def _extract_xlsx(path: Path) -> str:
    from initrunner._compat import require_ingest

    require_ingest("openpyxl")
    import openpyxl  # type: ignore[unresolved-import]

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in wb.worksheets:
            buf = io.StringIO()
            writer = csv.writer(buf)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow([str(c) if c is not None else "" for c in row])
            parts.append(f"# {sheet.title}\n{buf.getvalue()}")
    finally:
        wb.close()
    return "\n\n".join(parts)
